from openpyxl import load_workbook
from openpyxl.chart import Reference, series, BarChart3D

wb = load_workbook("test.xlsx")
ws = wb.active

data = Reference(ws, min_col=1, min_row=2, max_col=3, max_row=4)
titles = Reference(ws, min_row=1, min_col=1, max_col=3)
chart = BarChart3D()
chart.title = "Daily Sales Chart"
chart.add_data(data=data, titles_from_data=True)
chart.set_categories(titles)

ws.add_chart(chart, "D1")
wb.save("test_2.xlsx")